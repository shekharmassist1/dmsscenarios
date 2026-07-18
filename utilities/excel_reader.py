from openpyxl import load_workbook
import pandas as pd

def get_login_data():
    wb = load_workbook("data/login_data.xlsx")
    sheet = wb.active

    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        username, password, expected = row

        data.append(
            (
                username or "",
                password or "",
                expected
            )
        )

    return data


def get_bill_data():

    file_path = "testdata/bill_data.xlsx"

    df = pd.read_excel(file_path)

    return df.to_dict(orient="records")